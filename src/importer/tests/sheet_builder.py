"""Generate a small .xlsx that mimics the club spreadsheet layout, with invented players only.

The default content (``fixture_sheets()``) reproduces every known data problem:
- ТУР2026_new: a date with the wrong year, a text date with year 0202, a date column without
  results, a blank numbered slot, a merged name cell, a player without games, and a tie
  (payouts 300, 300, 100) whose cached place counts are LARGE-style (sheet semantics).
- ТУР2025: marker in column D, "Имя (ник)" style names (aliases of the ТУР2026 players), an
  orphan net formula (sheet counts a game that was not played), a wrong cached Итого, and one
  game whose payouts are 10 short of its buy-ins (leftover rule).
- КЭШ2026: cash sheet, text date, a fund block below the marker that must be ignored.
- ТУР2026_old version: must be skipped.

Summary cells hold plain numbers: they stand in for the cached formula values that openpyxl
reads back with data_only=True. Net cells hold formulas, as in the real sheet.
"""

import datetime
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl
from openpyxl.utils import column_index_from_string, get_column_letter


@dataclass
class FakeRow:
    name: str | None  # None: a numbered slot without a player
    results: dict[int, tuple[int, int]] = field(default_factory=dict)  # game -> (buyin, payout)
    cached: dict[str, int] = field(default_factory=dict)  # summary header -> cached value
    orphan_net: tuple[int, ...] = ()  # games with a net formula but no buy-in
    missing_net: tuple[int, ...] = ()  # played games without a net formula
    merge_name: bool = False


@dataclass
class FakeSheet:
    title: str
    start: str
    name_header: str
    summary: list[str]
    dates: list
    rows: list[FakeRow]
    marker_column: str = "B"
    below_marker: list[tuple[str, object]] = field(default_factory=list)  # (cell offset col, value)


TOUR_2026_SUMMARY = ["Итого", "Кол. игр", "ITM", "Кол. 1", "Кол. 2", "Кол. 3"]


def fixture_sheets() -> list[FakeSheet]:
    return [
        FakeSheet(
            title="ТУР2026_new",
            start="I",
            name_header="Имя - ник",
            summary=TOUR_2026_SUMMARY,
            dates=[
                datetime.datetime(2026, 1, 18),
                datetime.datetime(2025, 2, 1),  # wrong year -> 01.02.2026
                "14.03.0202",  # text date -> 14.03.2026
                "21.03.2026",  # no results: no game
            ],
            rows=[
                # g0: 300 in, 300 out; g1: 300/300; g2 (tie 300, 300, 100): 700/700
                FakeRow(
                    "Иван - Трактор",
                    {0: (100, 200), 2: (200, 300)},
                    tour_cached(net=200, games=2, itm=3, places=(2, 1, 0)),
                ),
                FakeRow(
                    "Пётр - Сеялка",
                    {0: (100, 100), 1: (100, 200), 2: (200, 300)},
                    tour_cached(net=200, games=3, itm=4, places=(2, 2, 0)),
                ),
                FakeRow(None),
                FakeRow(
                    "Мария - Комбайн",
                    {0: (100, 0), 1: (100, 100), 2: (200, 100)},
                    tour_cached(net=-200, games=3, itm=2, places=(0, 1, 1)),
                    merge_name=True,
                ),
                FakeRow(
                    "Олег - Плуг",
                    {1: (100, 0), 2: (100, 0)},
                    tour_cached(net=-200, games=2, itm=0, places=(0, 0, 0)),
                ),
                FakeRow("Нина - Борона", {}, tour_cached(net=0, games=0, itm=0, places=(0, 0, 0))),
            ],
        ),
        FakeSheet(
            title="ТУР2025",
            start="H",
            name_header="Имя (ник)",
            summary=["Итого", "Кол. игр", "Кол. 1", "Кол. 2", "Кол. 3"],
            dates=[
                datetime.datetime(2025, 5, 4),
                datetime.datetime(2025, 5, 11),
                datetime.datetime(2025, 5, 18),
            ],
            marker_column="D",
            rows=[
                # g0: 100/100; g1: 150 in, 140 out (10 short); g2: 100/100
                FakeRow(
                    "Иван (Трактор)",
                    {0: (50, 100), 1: (50, 90)},
                    {"Итого": 90, "Кол. игр": 2, "Кол. 1": 2, "Кол. 2": 0, "Кол. 3": 0},
                ),
                FakeRow(
                    "Пётр (Сеялка)",
                    {0: (50, 0), 1: (50, 50), 2: (50, 100)},
                    # Итого should be 0: a real mismatch
                    {"Итого": 20, "Кол. игр": 3, "Кол. 1": 1, "Кол. 2": 1, "Кол. 3": 0},
                ),
                FakeRow(
                    "Олег (Плуг)",
                    {1: (50, 0), 2: (50, 0)},
                    # the orphan net formula in g0 makes the sheet count 3 games
                    {"Итого": -100, "Кол. игр": 3, "Кол. 1": 0, "Кол. 2": 0, "Кол. 3": 0},
                    orphan_net=(0,),
                ),
            ],
        ),
        FakeSheet(
            title="КЭШ2026",
            start="E",
            name_header="Имя - ник",
            summary=["Итого", "Кол-во игр"],
            dates=[datetime.datetime(2026, 1, 18), "25.01.0202"],
            marker_column="D",
            below_marker=[("D", "Остаток"), ("D", "Закупка"), ("E", 10.0)],
            rows=[
                # g0: 250 in, 190 out (leftover 60); g1: 200 in, 195 out (leftover 5)
                FakeRow("Иван - Трактор", {0: (100, 150)}, {"Итого": 50, "Кол-во игр": 1}),
                FakeRow(
                    "Мария - Комбайн",
                    {0: (100, 40), 1: (100, 95)},
                    {"Итого": -65, "Кол-во игр": 2},
                ),
                FakeRow(
                    "Олег - Плуг", {0: (50, 0), 1: (100, 100)}, {"Итого": -50, "Кол-во игр": 2}
                ),
            ],
        ),
        FakeSheet(
            title="ТУР2026_old version",
            start="I",
            name_header="Имя - ник",
            summary=TOUR_2026_SUMMARY,
            dates=[datetime.datetime(2026, 1, 18)],
            rows=[FakeRow("Старый Формат", {0: (100, 100)})],
        ),
    ]


def tour_cached(*, net: int, games: int, itm: int, places: tuple[int, int, int]) -> dict:
    return {
        "Итого": net,
        "Кол. игр": games,
        "ITM": itm,
        "Кол. 1": places[0],
        "Кол. 2": places[1],
        "Кол. 3": places[2],
    }


def fixture_aliases() -> dict:
    """aliases.yaml content for fixture_sheets(); exercises ё/е, case and space matching."""
    return {
        "players": [
            {
                "name": "Иван Тестов",
                "nickname": "Трактор",
                "aliases": ["Иван - Трактор", "Иван (Трактор)"],
            },
            {
                "name": "Пётр Примеров",
                "nickname": "Сеялка",
                "aliases": ["Петр - Сеялка", "Пётр (Сеялка)"],
            },
            {"name": "Мария Образцова", "nickname": "Комбайн", "aliases": ["Мария - Комбайн"]},
            {"name": "Олег Пробный", "slug": "oleg", "aliases": ["Олег - Плуг", "олег   (ПЛУГ)"]},
        ]
    }


def build_workbook(path: Path, sheets: list[FakeSheet] | None = None) -> Path:
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    for sheet in sheets if sheets is not None else fixture_sheets():
        _write_sheet(workbook.create_sheet(sheet.title), sheet)
    workbook.save(path)
    return path


def sheet_named(sheets: list[FakeSheet], title: str) -> FakeSheet:
    """Pick one sheet to edit; fixture_sheets() returns fresh objects on every call."""
    return next(sheet for sheet in sheets if sheet.title == title)


def _write_sheet(ws, sheet: FakeSheet) -> None:
    start = column_index_from_string(sheet.start)
    ws["A1"], ws["B1"] = "№", sheet.name_header
    for offset, header in enumerate(sheet.summary):
        ws.cell(1, 3 + offset, header)
    for index, value in enumerate(sheet.dates):
        column = start + 2 * index
        ws.cell(1, column, value)
        ws.merge_cells(start_row=1, start_column=column, end_row=1, end_column=column + 1)

    row = 2
    for number, fake in enumerate(sheet.rows, start=1):
        ws.cell(row, 1, number)
        if fake.name is not None:
            ws.cell(row, 2, fake.name)
            if fake.merge_name:
                ws.merge_cells(start_row=row, start_column=2, end_row=row + 1, end_column=2)
            for offset, header in enumerate(sheet.summary):
                if header in fake.cached:
                    ws.cell(row, 3 + offset, fake.cached[header])
            for game in range(len(sheet.dates)):
                column = start + 2 * game
                letter = get_column_letter(column)
                net_formula = f"={letter}{row + 1}-{letter}{row}"
                if game in fake.results:
                    buyin, payout = fake.results[game]
                    ws.cell(row, column, float(buyin))
                    ws.cell(row + 1, column, float(payout))
                    if game not in fake.missing_net:
                        ws.cell(row, column + 1, net_formula)
                elif game in fake.orphan_net:
                    ws.cell(row, column + 1, net_formula)
        row += 2

    marker = column_index_from_string(sheet.marker_column)
    ws.cell(row, marker, "Закупка")
    for offset, (letter, value) in enumerate(sheet.below_marker, start=1):
        ws.cell(row + offset, column_index_from_string(letter), value)
