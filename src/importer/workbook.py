"""Read the club spreadsheet into plain dataclasses. No database access here.

Layout (the same on every import sheet):
- Row 1: headers. Summary columns (``Итого``, ``Кол. игр`` ...) come first, then game dates in
  every second column; the column after each date holds a net formula and is not data.
- Each player takes two rows: the name row (column B, buy-in under the date) and the payout row.
- Player rows end at the first row where column B or D reads ``Закупка``.
"""

import datetime
import io
import re
import zipfile
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.worksheet.worksheet import Worksheet

from club.models import SeasonKind

SHEET_NAME_RE = re.compile(r"^(ТУР|КЭШ)(\d{4})(_new)?$")
KIND_BY_PREFIX = {"ТУР": SeasonKind.TOUR, "КЭШ": SeasonKind.CASH}
TEXT_DATE_RE = re.compile(r"^\s*(\d{1,2})\.(\d{1,2})\.(\d{2,4})\s*$")
END_MARKER = "Закупка"
END_MARKER_COLUMNS = (2, 4)  # B, D (ТУР2025 has it in D)
NAME_COLUMN = 2
FIRST_DATA_ROW = 2
# Cached summary columns, found by header text. Missing ones are simply not verified.
SUMMARY_HEADERS = {
    "net": ("Итого",),
    "games": ("Кол. игр", "Кол-во игр"),
    "itm": ("ITM",),
    "first": ("Кол. 1",),
    "second": ("Кол. 2",),
    "third": ("Кол. 3",),
}


class SheetError(Exception):
    """One or more problems in the workbook; ``errors`` lists them all with cell references."""

    def __init__(self, errors: list[str]):
        self.errors = errors
        super().__init__("\n".join(errors))


@dataclass(frozen=True)
class DateFix:
    cell: str
    raw: str
    fixed: datetime.date
    reason: str

    def __str__(self) -> str:
        return f"{self.cell}: {self.raw} -> {self.fixed:%d.%m.%Y} ({self.reason})"


@dataclass(frozen=True)
class GameColumn:
    letter: str
    date: datetime.date


@dataclass(frozen=True)
class Entry:
    buyin: int
    payout: int


@dataclass
class PlayerRow:
    row: int
    raw_name: str
    entries: dict[datetime.date, Entry] = field(default_factory=dict)
    cached: dict[str, int | None] = field(default_factory=dict)
    # Sheet-side formula gaps, used only to explain verification differences.
    orphan_net: dict[datetime.date, int] = field(default_factory=dict)  # date -> formula value
    missing_net: list[datetime.date] = field(default_factory=list)


@dataclass
class ParsedSheet:
    title: str
    year: int
    kind: str
    start_letter: str
    games: list[GameColumn]
    players: list[PlayerRow]
    summary_columns: dict[str, str]
    date_fixes: list[DateFix] = field(default_factory=list)
    empty_columns: list[str] = field(default_factory=list)


@dataclass
class ParsedWorkbook:
    sheets: list[ParsedSheet]
    skipped: list[str]


def season_of(title: str) -> tuple[int, str] | None:
    match = SHEET_NAME_RE.match(title)
    if not match:
        return None
    return int(match.group(2)), KIND_BY_PREFIX[match.group(1)]


def clean_name(value: str) -> str:
    return " ".join(value.split())


def parse_workbook(source: str | Path | bytes, only: list[str] | None = None) -> ParsedWorkbook:
    """Parse every import sheet (or only the named ones); raise SheetError listing all problems.

    ``source`` is a path or the file's bytes (the sheet sync keeps fetched bytes, never a file).
    """
    formulas = _load(source)
    cached = _load(source, data_only=True)
    errors: list[str] = []

    if only is not None:
        unknown = [name for name in only if name not in formulas.sheetnames]
        errors += [f"{name}: no such sheet" for name in unknown]
        errors += [
            f"{name}: not an import sheet (expected ТУРyyyy or КЭШyyyy)"
            for name in only
            if name not in unknown and season_of(name) is None
        ]

    sheets, skipped, seen = [], [], {}
    for ws in formulas.worksheets:
        season = season_of(ws.title)
        if season is None:
            skipped.append(ws.title)
            continue
        if only is not None and ws.title not in only:
            continue
        if season in seen:
            errors.append(f"{ws.title}: same season as {seen[season]}")
            continue
        seen[season] = ws.title
        try:
            sheets.append(_parse_sheet(ws, cached[ws.title], *season))
        except SheetError as exc:
            errors += exc.errors

    if errors:
        raise SheetError(errors)
    return ParsedWorkbook(sheets=sheets, skipped=skipped)


def _load(source: str | Path | bytes, **options) -> openpyxl.Workbook:
    try:
        return openpyxl.load_workbook(
            io.BytesIO(source) if isinstance(source, bytes) else source, **options
        )
    except (zipfile.BadZipFile, InvalidFileException, KeyError) as exc:
        raise SheetError([f"not a readable .xlsx workbook ({type(exc).__name__})"]) from exc


def _parse_sheet(ws: Worksheet, cached_ws: Worksheet, year: int, kind: str) -> ParsedSheet:
    errors: list[str] = []

    def ref(row: int, column: int) -> str:
        return f"{ws.title}!{get_column_letter(column)}{row}"

    # Game columns: from the first date-like header, every second column while dates continue.
    start = next(
        (c for c in range(3, ws.max_column + 1) if _is_date_like(ws.cell(1, c).value)), None
    )
    if start is None:
        raise SheetError([f"{ws.title}: no date columns in row 1"])
    columns: list[tuple[int, datetime.date]] = []
    date_fixes: list[DateFix] = []
    column = start
    while column <= ws.max_column and _is_date_like(ws.cell(1, column).value):
        try:
            date, fix = _fix_date(ws.cell(1, column).value, year, ref(1, column))
        except ValueError as exc:
            errors.append(str(exc))
        else:
            columns.append((column, date))
            if fix:
                date_fixes.append(fix)
        column += 2
    for (col_a, date), (col_b, other) in _pairs(columns):
        if date == other:
            errors.append(f"{ref(1, col_b)}: date {date:%d.%m.%Y} repeats {ref(1, col_a)}")

    summary_columns = {}
    for key, labels in SUMMARY_HEADERS.items():
        for c in range(3, start):
            if isinstance(ws.cell(1, c).value, str) and ws.cell(1, c).value.strip() in labels:
                summary_columns[key] = c

    end = next(
        (
            r
            for r in range(FIRST_DATA_ROW, ws.max_row + 1)
            if any(_text(ws.cell(r, c).value) == END_MARKER for c in END_MARKER_COLUMNS)
        ),
        None,
    )
    if end is None:
        raise SheetError(errors + [f"{ws.title}: end marker '{END_MARKER}' not found in B or D"])

    players = []
    for row in range(FIRST_DATA_ROW, end, 2):
        name = ws.cell(row, NAME_COLUMN).value
        if name is not None and not isinstance(name, str):
            errors.append(f"{ref(row, NAME_COLUMN)}: name is not text")
            continue
        name = clean_name(name or "")
        if not name:
            for c, _ in columns:
                for r in (row, row + 1):
                    if ws.cell(r, c).value is not None:
                        errors.append(f"{ref(r, c)}: number in a row without a player name")
            continue

        player = PlayerRow(row=row, raw_name=name)
        for key, c in summary_columns.items():
            player.cached[key] = _cached_int(cached_ws.cell(row, c).value)
        for c, date in columns:
            buyin_cell, payout_cell = ws.cell(row, c).value, ws.cell(row + 1, c).value
            net_cell = ws.cell(row, c + 1).value
            has_net_formula = isinstance(net_cell, str) and net_cell.startswith("=")
            try:
                payout = _amount(payout_cell, ref(row + 1, c)) or 0
                buyin = _amount(buyin_cell, ref(row, c))
            except ValueError as exc:
                errors.append(str(exc))
                continue
            if buyin is None:
                if payout:
                    errors.append(f"{ref(row + 1, c)}: payout without a buy-in")
                if has_net_formula:
                    player.orphan_net[date] = payout
                continue
            if buyin == 0:
                errors.append(f"{ref(row, c)}: buy-in is 0 (leave it empty if they did not play)")
                continue
            player.entries[date] = Entry(buyin=buyin, payout=payout)
            if not has_net_formula:
                player.missing_net.append(date)
        players.append(player)

    if errors:
        raise SheetError(errors)

    played = {date for player in players for date in player.entries}
    return ParsedSheet(
        title=ws.title,
        year=year,
        kind=kind,
        start_letter=get_column_letter(start),
        games=[GameColumn(get_column_letter(c), date) for c, date in columns if date in played],
        players=players,
        summary_columns={key: get_column_letter(c) for key, c in summary_columns.items()},
        date_fixes=date_fixes,
        empty_columns=[
            f"{get_column_letter(c)} ({date:%d.%m.%Y})" for c, date in columns if date not in played
        ],
    )


def _pairs(items):
    return [(a, b) for i, a in enumerate(items) for b in items[i + 1 :]]


def _text(value) -> str:
    return value.strip() if isinstance(value, str) else ""


def _is_date_like(value) -> bool:
    return isinstance(value, datetime.datetime | datetime.date) or bool(
        isinstance(value, str) and TEXT_DATE_RE.match(value)
    )


def _fix_date(value, year: int, cell: str) -> tuple[datetime.date, DateFix | None]:
    """Take day and month from the cell and the year from the sheet name."""
    if isinstance(value, datetime.datetime | datetime.date):
        day, month, raw_year, raw = value.day, value.month, value.year, f"{value:%d.%m.%Y}"
        reason = "wrong year"
    else:
        match = TEXT_DATE_RE.match(value)
        day, month, raw_year = int(match.group(1)), int(match.group(2)), int(match.group(3))
        raw, reason = f'"{value.strip()}"', "text date"
    try:
        fixed = datetime.date(year, month, day)
    except ValueError:
        raise ValueError(f"{cell}: {raw} is not a valid date") from None
    if raw_year == year:
        return fixed, None
    return fixed, DateFix(cell=cell, raw=raw, fixed=fixed, reason=reason)


def _amount(value, cell: str) -> int | None:
    """Integer lei from a cell, None when empty."""
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    if isinstance(value, bool) or not isinstance(value, int | float):
        raise ValueError(f"{cell}: expected a number, got {value!r}")
    if value != int(value):
        raise ValueError(f"{cell}: {value} is not a whole number of lei")
    if value < 0:
        raise ValueError(f"{cell}: negative amount {value}")
    return int(value)


def _cached_int(value) -> int | None:
    if isinstance(value, bool) or not isinstance(value, int | float):
        return None
    return round(value)
